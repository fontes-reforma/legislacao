#!/usr/bin/env python3
"""
Gera um site estático com a legislação da reforma tributária fatiada por
artigo (uma página por artigo/anexo), para uso como fonte de conhecimento
de agentes que buscam via Bing (Copilot).

Uso:
    SITE_BASE_URL=https://usuario.github.io/repo python scripts/build_site.py

Saída: pasta _site/ (HTML + sitemap.xml + robots.txt + chave IndexNow).
"""
from __future__ import annotations

import datetime as dt
import html
import json
import os
import re
import sys
import unicodedata
from pathlib import Path

import requests
from bs4 import BeautifulSoup

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "_site"
MAX_CHARS = 12000          # páginas maiores que isso são divididas em partes
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0 Safari/537.36")
HOJE = dt.date.today()

avisos: list[str] = []


GHA = os.environ.get("GITHUB_ACTIONS") == "true"


def aviso(msg: str) -> None:
    avisos.append(msg)
    print(f"[AVISO] {msg}", file=sys.stderr)
    if GHA:
        print(f"::warning title=Aviso do build::{msg}")


def nota(titulo: str, msg: str) -> None:
    """Publica um resumo como anotação do GitHub Actions (legível pela API)."""
    if GHA:
        msg = msg.replace("%", "%25").replace("\r", "%0D").replace("\n", "%0A")
        print(f"::notice title={titulo}::{msg}")


# --------------------------------------------------------------------------
# Download
# --------------------------------------------------------------------------
CACHE = RAIZ / "cache"


def obter_bytes(fonte: dict) -> bytes | None:
    """Ordem: arquivo manual > download (3 tentativas) > última versão boa em cache/.

    Um download bem-sucedido atualiza o cache, exceto se vier com menos da
    metade do tamanho da versão em cache (provável página de erro).
    """
    import time

    manual = RAIZ / fonte.get("arquivo_manual", "")
    if fonte.get("arquivo_manual") and manual.is_file():
        print(f"[{fonte['id']}] usando arquivo manual {manual.name}")
        return manual.read_bytes()

    cache = CACHE / f"{fonte['id']}.bin"
    anterior = cache.read_bytes() if cache.is_file() else None
    url = fonte.get("url")
    erro = "sem URL configurada"
    if url:
        for tentativa in range(3):
            try:
                r = requests.get(url, headers={"User-Agent": UA}, timeout=120)
                r.raise_for_status()
                if anterior and len(r.content) < len(anterior) / 2:
                    erro = (f"download suspeito ({len(r.content):,} bytes contra "
                            f"{len(anterior):,} do cache)")
                    break
                CACHE.mkdir(exist_ok=True)
                cache.write_bytes(r.content)
                print(f"[{fonte['id']}] baixado {len(r.content):,} bytes de {url}")
                return r.content
            except Exception as e:  # noqa: BLE001
                erro = str(e)
                time.sleep(10 * (tentativa + 1))

    if anterior:
        aviso(f"{fonte['id']}: {erro}; usando a última versão baixada com sucesso (cache)")
        return anterior
    aviso(f"{fonte['id']}: {erro}; sem cache. Salve o arquivo em {fonte.get('arquivo_manual')}")
    return None


# --------------------------------------------------------------------------
# Extração de parágrafos
# --------------------------------------------------------------------------
def limpar(txt: str) -> str:
    txt = txt.replace("\xa0", " ")
    return re.sub(r"\s+", " ", txt).strip()


def decodificar(raw: bytes) -> str:
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def paragrafos_planalto(raw: bytes) -> list[str]:
    """Extrai parágrafos do HTML do Planalto, SEM o texto tachado
    (redações revogadas/alteradas, que o Planalto mantém riscadas)."""
    soup = BeautifulSoup(decodificar(raw), "lxml")
    for tag in soup.find_all(["strike", "s", "del", "script", "style"]):
        tag.decompose()
    # o Planalto também risca via CSS em alguns textos
    for tag in soup.find_all(style=re.compile(r"line-through", re.I)):
        tag.decompose()

    saida: list[str] = []
    for el in soup.find_all(["p", "tr"]):
        if el.name == "p" and el.find_parent("tr") is not None:
            continue  # conteúdo de tabela é tratado por linha
        if el.name == "tr":
            celulas = [limpar(td.get_text(" ")) for td in el.find_all(["td", "th"])]
            txt = " | ".join(c for c in celulas if c)
        else:
            txt = limpar(el.get_text(" "))
        if txt:
            saida.append(txt)
    return saida


INICIO_BLOCO = re.compile(
    r"^(Art\.\s*\d|§\s*\d|Parágrafo único|[IVXLCDM]+\s*[-–]\s|[a-z]\)\s|"
    r"LIVRO\s|TÍTULO\s|CAPÍTULO\s|Seção\s|SEÇÃO\s|Subseção\s|SUBSEÇÃO\s|ANEXO\s)"
)


def linhas_pdf(raw: bytes) -> list[str]:
    import io
    import pdfplumber

    paginas: list[list[str]] = []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for pg in pdf.pages:
            t = pg.extract_text() or ""
            paginas.append([limpar(l) for l in t.splitlines() if limpar(l)])

    # remove cabeçalhos/rodapés que se repetem em mais da metade das páginas
    from collections import Counter
    cont = Counter(l for pg in paginas for l in set(pg))
    n = max(len(paginas), 1)
    repetidas = {l for l, c in cont.items() if n >= 4 and c > n / 2}
    linhas = [l for pg in paginas for l in pg if l not in repetidas
              and not re.fullmatch(r"(p[áa]g(ina)?\.?\s*)?\d+(\s*(de|/)\s*\d+)?", l, re.I)]
    return linhas


def paragrafos_pdf(raw: bytes) -> list[str]:
    paras: list[str] = []
    for l in linhas_pdf(raw):
        if not paras or INICIO_BLOCO.match(l):
            paras.append(l)
        elif paras[-1].endswith("-") and not paras[-1].endswith(" -"):
            paras[-1] = paras[-1][:-1] + l
        else:
            paras[-1] += " " + l
    return paras


# --------------------------------------------------------------------------
# Segmentação em unidades (artigos, anexos)
# --------------------------------------------------------------------------
RE_ART = re.compile(
    r"^Art\.\s*(\d{1,4})\s*(?:º|°|o)?\s*(?:-\s*([A-Z]{1,2}))?\s*[\.\-–]?(?=\s|$)")
RE_CAB = re.compile(
    r"^(LIVRO|TÍTULO|CAPÍTULO|SEÇÃO|Seção|SUBSEÇÃO|Subseção)\s+"
    r"([IVXLCDM]+(?:-[A-Z])?|ÚNIC[AO]|Únic[ao])\b\s*[-–—.]?\s*(.*)$")
RE_ANEXO = re.compile(r"^ANEXO\s+([IVXLCDM]+(?:-[A-Z])?)\b\s*[-–—.]?\s*(.*)$")
RE_ADCT = re.compile(r"^ATO DAS DISPOSI[ÇC][ÕO]ES CONSTITUCIONAIS TRANSIT[ÓO]RIAS\s*$", re.I)
RE_CLAUSULA = re.compile(
    r"^Cl[áa]usula\s+((?:d[ée]cima|vig[ée]sima|trig[ée]sima)(?:\s+\w+)?|\w+)(?=[\s.:–-])", re.I)
_UNID = {"primeira": 1, "segunda": 2, "terceira": 3, "quarta": 4, "quinta": 5, "sexta": 6,
         "sétima": 7, "setima": 7, "oitava": 8, "nona": 9}
_DEZ = {"décima": 10, "decima": 10, "vigésima": 20, "vigesima": 20, "trigésima": 30, "trigesima": 30}


def ordinal(txt: str) -> int:
    partes = txt.lower().split()
    if len(partes) == 1:
        return _UNID.get(partes[0]) or _DEZ.get(partes[0], 0)
    return _DEZ.get(partes[0], 0) + _UNID.get(partes[1], 0) if partes[0] in _DEZ else 0


RE_PROD = re.compile(r"\s*Produ[çc][ãa]o de efeitos\s*", re.I)
NIVEIS = ["livro", "titulo", "capitulo", "secao", "subsecao"]


def nivel_de(palavra: str) -> str:
    p = unicodedata.normalize("NFKD", palavra.upper()).encode("ascii", "ignore").decode()
    return {"LIVRO": "livro", "TITULO": "titulo", "CAPITULO": "capitulo",
            "SECAO": "secao", "SUBSECAO": "subsecao"}[p]


def segmentar(paras: list[str]) -> list[dict]:
    unidades: list[dict] = []
    cab: dict[str, str] = {}
    pendente: str | None = None       # nível aguardando o nome na próxima linha
    atual: dict | None = None
    ultimo = (0, "")
    ultima_clausula = 0
    em_anexos = False
    prefixo = ""          # "adct" após o título do ADCT (CF), que reinicia a numeração

    def nova(tipo, rotulo, slug):
        u = {"tipo": tipo, "rotulo": rotulo, "slug": slug,
             "contexto": dict(cab), "paras": []}
        unidades.append(u)
        return u

    for p in paras:
        m_anx = RE_ANEXO.match(p)
        if m_anx and (ultimo[0] > 0 or em_anexos):
            em_anexos = True
            cab.clear()
            rom = m_anx.group(1)
            # Rótulo duplo: "ANEXO XVIII" + "(Lei Complementar nº 123...)" + "ANEXO I".
            # O segundo "ANEXO" é a numeração na lei alterada, não um anexo novo.
            if (atual is not None and atual["tipo"] == "anexo" and len(atual["paras"]) <= 3
                    and not any(" | " in x for x in atual["paras"])):
                lei = next((re.search(r"Lei Complementar n[ºo°]\s*([\d.]+).*?(\d{4})\)?\s*$", x, re.I)
                            for x in atual["paras"] if "Lei Complementar" in x), None)
                ref = f" da LC {lei.group(1)}/{lei.group(2)}" if lei else ""
                atual["rotulo"] += f" (Anexo {rom}{ref})"
                atual["paras"].append(p)
                atual["nome"] = RE_PROD.sub("", m_anx.group(2)).strip()
                atual["pegar_nome"] = not atual["nome"]
                pendente = None
                continue
            atual = nova("anexo", f"Anexo {rom}", f"anexo-{rom.lower()}")
            atual["nome"] = RE_PROD.sub("", m_anx.group(2)).strip()
            atual["pegar_nome"] = not atual["nome"]
            atual["paras"].append(p)
            pendente = None
            continue
        if atual is not None and atual.get("pegar_nome"):
            atual["nome"] = p[:160]
            atual["pegar_nome"] = False

        m_cab = RE_CAB.match(p)
        if m_cab and not em_anexos:
            nv = nivel_de(m_cab.group(1))
            # zera os níveis inferiores
            for n in NIVEIS[NIVEIS.index(nv):]:
                cab.pop(n, None)
            ident = f"{m_cab.group(1).capitalize()} {m_cab.group(2)}"
            nome = m_cab.group(3).strip()
            cab[nv] = f"{ident} – {nome}" if nome else ident
            pendente = None if nome else nv
            continue

        if RE_ADCT.match(p) and not em_anexos:
            prefixo, ultimo, pendente = "adct", (0, ""), None
            cab.clear()
            cab["livro"] = "Ato das Disposições Constitucionais Transitórias"
            continue

        m_cl = RE_CLAUSULA.match(p)
        if m_cl and not em_anexos:
            n = ordinal(m_cl.group(1))
            if n and n > ultima_clausula:
                ultima_clausula = n
                atual = nova("artigo", f"Cláusula {m_cl.group(1).lower()}", f"clausula-{n}")
                atual["paras"].append(p)
                pendente = None
                continue

        m_art = RE_ART.match(p)
        if m_art and not em_anexos:
            chave = (int(m_art.group(1)), m_art.group(2) or "")
            if chave > ultimo:
                ultimo = chave
                num, suf = chave
                rot = f"{num}{'º' if num < 10 else ''}{'-' + suf if suf else ''}"
                slug = f"art-{num}{'-' + suf.lower() if suf else ''}"
                rotulo = f"Art. {rot}"
                if prefixo:
                    slug, rotulo = f"{prefixo}-{slug}", f"ADCT, {rotulo}"
                atual = nova("artigo", rotulo, slug)
                atual["paras"].append(p)
                pendente = None
                continue

        if pendente and len(p) < 200 and not INICIO_BLOCO.match(p):
            cab[pendente] = f"{cab[pendente]} – {p}"
            pendente = None
            continue
        pendente = None

        if atual is None:
            atual = nova("preambulo", "Ementa e preâmbulo", "preambulo")
        atual["paras"].append(p)

    return unidades


RE_SECAO = re.compile(r"^(\d{1,2}(?:\.\d{1,2}){0,3})\.?\s+[A-ZÁÉÍÓÚÂÊÔÃÕÇ][^|]{2,110}$")
RE_ITEM = re.compile(r"^([•▪\-–o]\s|\(?[a-z]\)\s|\d{3}[a-z]?\s)")


def paragrafos_nt(raw: bytes) -> list[str]:
    """Parágrafos de notas técnicas: quebra em títulos de seção, itens e fim de frase."""
    paras: list[str] = []
    for l in linhas_pdf(raw):
        if re.match(r"^[.\s]{0,3}$", l) or re.search(r"\.{6,}\s*\d+$", l):
            continue  # linhas do sumário
        novo = (not paras or RE_SECAO.match(l) or RE_ITEM.match(l)
                or paras[-1].endswith((".", ":", ";")) or bool(RE_SECAO.match(paras[-1])))
        if novo:
            paras.append(l)
        elif paras[-1].endswith("-") and not paras[-1].endswith(" -"):
            paras[-1] = paras[-1][:-1] + l
        else:
            paras[-1] += " " + l
    out = []
    for p in paras:  # parágrafos gigantes (tabelas) viram pedaços de ~1.500 caracteres
        while len(p) > 1800:
            corte = max(p.rfind(". ", 0, 1500), p.rfind(" ", 0, 1500))
            corte = corte if corte > 500 else 1500
            out.append(p[:corte + 1].strip())
            p = p[corte + 1:].strip()
        out.append(p)
    return out


def segmentar_blocos(paras: list[str]) -> list[dict]:
    """Documentos sem artigos (notas técnicas): blocos por seção, com ~4.000 caracteres."""
    blocos, buf, tam, titulo = [], [], 0, ""
    for p in paras:
        secao = RE_SECAO.match(p)
        if buf and ((secao and tam > 1200) or tam + len(p) > 4000):
            blocos.append((titulo, buf))
            buf, tam = [], 0
        if secao and (not buf or not titulo):
            titulo = p[:120]
        elif not buf and not titulo:
            titulo = p[:120]
        if secao and not buf:
            titulo = p[:120]
        buf.append(p)
        tam += len(p)
    if buf:
        blocos.append((titulo, buf))
    out = []
    for i, (t, b) in enumerate(blocos, 1):
        out.append({"tipo": "bloco", "rotulo": f"Parte {i}", "slug": f"parte-{i}",
                    "nome": t, "contexto": {}, "paras": b})
    return out


# --------------------------------------------------------------------------
# Tabelas (cClassTrib, cIndOp, correlação NBS)
# --------------------------------------------------------------------------
def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = limpar(str(v))
    return re.sub(r"^(\d{4}-\d{2}-\d{2}) 00:00:00$", r"\1", s)


def _linhas_xlsx(raw: bytes, aba_contem: str):
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    for ws in wb.worksheets:
        linhas = [[_txt(c) for c in r] for r in ws.iter_rows(values_only=True)]
        linhas = [l for l in linhas if any(l)]
        if linhas and any(aba_contem.lower() == c.lower() for c in linhas[0]):
            return linhas
    return []


def unidades_cclasstrib(raw: bytes) -> list[dict]:
    cst = _linhas_xlsx(raw, "Descrição CST-IBS/CBS")
    cct = _linhas_xlsx(raw, "cClassTrib")
    # a aba CST também tem "Descrição CST-IBS/CBS"; a de cClassTrib tem a coluna cClassTrib
    if cst and "cClassTrib" in cst[0]:
        cst, cct = [], cst
    unidades: list[dict] = []
    if not cct:
        raise ValueError("aba cClassTrib não encontrada")
    h = cct[0]
    ix = {n: i for i, n in enumerate(h)}
    docs = [c for c in h if c.startswith("ind") and c[3:4].isupper() and c not in ("indDERE",)]
    por_cst: dict[str, list] = {}
    for r in cct[1:]:
        r = r + [""] * (len(h) - len(r))
        g = lambda n: r[ix[n]] if n in ix else ""
        cod = g("cClassTrib").zfill(6)
        cst_cod = g("CST-IBS/CBS").zfill(3)
        disp = g("LC 214/25")
        nome = g("Nome cClassTrib")
        por_cst.setdefault(cst_cod, []).append((cod, nome, disp))
        ps = [f"cClassTrib {cod}: {nome}.",
              f"CST-IBS/CBS {cst_cod}: {g('Descrição CST-IBS/CBS')}.",
              f"Descrição: {g('Descrição cClassTrib')}"]
        if disp:
            ps.append(f"Dispositivo legal: LC 214/2025, {disp}.")
        ps.append(f"Tipo de alíquota: {g('Tipo de Alíquota')}. Redução do IBS: {g('pRedIBS') or '0'}%. "
                  f"Redução da CBS: {g('pRedCBS') or '0'}%.")
        inds = [f"{n} = {g(n)}" for n in h if n.startswith("ind_") and g(n) not in ("", "0", "N/A")]
        if inds:
            ps.append("Indicadores ativos: " + "; ".join(inds) + ".")
        if g("Crédito para"):
            ps.append(f"Crédito para: {g('Crédito para')}.")
        permitidos = [d[3:] for d in docs if g(d).strip().lower() in ("1", "sim", "s")]
        if permitidos:
            ps.append("Documentos fiscais em que o código é aceito: " + ", ".join(permitidos) + ".")
        vig = " a ".join(x for x in (g("dIniVig"), g("dFimVig")) if x)
        if vig:
            ps.append(f"Vigência: {vig}.")
        if g("DataAtualização"):
            ps.append(f"Data da última atualização na tabela: {g('DataAtualização')}.")
        if g("ANEXO"):
            ps.append(f"Anexo da LC 214/2025: {g('ANEXO')}.")
        if g("LC Redação"):
            ps.append(f"Texto do dispositivo: {g('LC Redação')}")
        unidades.append({"tipo": "codigo", "rotulo": f"cClassTrib {cod}", "slug": f"cclasstrib-{cod}",
                         "nome": f"{nome} (LC 214/2025, {disp})" if disp else nome,
                         "contexto": {"livro": f"CST {cst_cod}"}, "paras": ps})
    # páginas de CST
    hc = cst[0] if cst else []
    ixc = {n: i for i, n in enumerate(hc)}
    desc_cst = {}
    for r in cst[1:]:
        r = r + [""] * (len(hc) - len(r))
        c = r[ixc["CST-IBS/CBS"]].zfill(3)
        inds = [f"{n} = {r[ixc[n]]}" for n in hc if n.startswith("ind") and r[ixc[n]] not in ("", "0")]
        desc_cst[c] = (r[ixc.get("Descrição CST-IBS/CBS", 1)], inds)
    cst_units = []
    for c in sorted(set(por_cst) | set(desc_cst)):
        desc, inds = desc_cst.get(c, ("", []))
        if not desc and por_cst.get(c):
            desc = next((r[ix["Descrição CST-IBS/CBS"]] for r in cct[1:]
                         if r[ix["CST-IBS/CBS"]].zfill(3) == c), "")
        ps = [f"CST-IBS/CBS {c}: {desc}."]
        if inds:
            ps.append("Indicadores do CST: " + "; ".join(inds) + ".")
        for cod, nome, disp in por_cst.get(c, []):
            ps.append(f"cClassTrib {cod}: {nome}" + (f" (LC 214/2025, {disp})" if disp else "") + ".")
        cst_units.append({"tipo": "codigo", "rotulo": f"CST {c}", "slug": f"cst-{c}",
                          "nome": desc, "contexto": {}, "paras": ps})
    return cst_units + unidades


def unidades_indop(raw: bytes) -> list[dict]:
    import csv
    import io
    linhas = [[limpar(c) for c in l] for l in csv.reader(io.StringIO(decodificar(raw)))]
    linhas = [l for l in linhas if any(l)]
    h = linhas[0]
    out = []
    for r in linhas[1:]:
        r = r + [""] * (len(h) - len(r))
        cod = r[0].zfill(6)
        ps = [f"Código indicador da operação (cIndOp) {cod}."]
        ps += [f"{h[i]}: {r[i]}." for i in range(1, len(h)) if r[i]]
        out.append({"tipo": "codigo", "rotulo": f"cIndOp {cod}", "slug": f"cindop-{cod}",
                    "nome": " – ".join(x for x in r[1:3] if x), "contexto": {}, "paras": ps})
    return out


def unidades_nbs(raw: bytes) -> list[dict]:
    linhas = _linhas_xlsx(raw, "Item LC 116")
    h = linhas[0]
    ix = {n: i for i, n in enumerate(h) if n}
    grupos: dict[str, dict] = {}
    ordem = []
    atual_item, herdado = "", {}
    herdaveis = ["PS ONEROSA? (S/N)", "ADQ EXTERIOR? (S/N)", "INDOP", "Local incidência IBS",
                 "cClassTrib", "nome cClassTrib"]
    for r in linhas[1:]:
        r = r + [""] * (len(h) - len(r))
        g = lambda n: r[ix[n]] if n in ix else ""
        if g("Item LC 116"):
            atual_item = g("Item LC 116")
            herdado = {}
            if atual_item not in grupos:
                grupos[atual_item] = {"desc": g("Descrição Item"), "linhas": []}
                ordem.append(atual_item)
        if not atual_item:
            continue
        for n in herdaveis:
            if g(n):
                herdado[n] = g(n)
        if g("NBS"):
            grupos[atual_item]["linhas"].append(
                f"NBS {g('NBS')} – {g('DESCRIÇÃO NBS')} | prestação onerosa: "
                f"{herdado.get('PS ONEROSA? (S/N)', '')} | adquirente no exterior: "
                f"{herdado.get('ADQ EXTERIOR? (S/N)', '')} | cIndOp: {herdado.get('INDOP', '')} | "
                f"local de incidência do IBS: {herdado.get('Local incidência IBS', '')} | "
                f"cClassTrib: {herdado.get('cClassTrib', '')} – {herdado.get('nome cClassTrib', '')}")
    out = []
    for it in ordem:
        gpo = grupos[it]
        out.append({"tipo": "codigo", "rotulo": f"Item {it} da LC 116/2003",
                    "slug": "item-" + re.sub(r"[^0-9a-z]+", "-", it.lower()).strip("-"),
                    "nome": gpo["desc"], "contexto": {},
                    "paras": [f"Item {it} da lista da LC 116/2003: {gpo['desc'].rstrip('.')}. "
                              "Correlação com NBS, código indicador da operação (cIndOp) e cClassTrib."]
                             + gpo["linhas"]})
    return out


def dividir_grandes(unidades: list[dict]) -> list[dict]:
    out = []
    for u in unidades:
        total = sum(len(p) for p in u["paras"])
        if total <= MAX_CHARS:
            out.append(u)
            continue
        partes, buf, tam = [], [], 0
        for p in u["paras"]:
            if tam + len(p) > MAX_CHARS and buf:
                partes.append(buf)
                buf, tam = [], 0
            buf.append(p)
            tam += len(p)
        partes.append(buf)
        for i, b in enumerate(partes, 1):
            v = dict(u)
            v["paras"] = b
            if i > 1:
                v["slug"] = f"{u['slug']}-parte-{i}"
                v["rotulo"] = f"{u['rotulo']} (parte {i} de {len(partes)})"
            else:
                v["rotulo"] = f"{u['rotulo']} (parte 1 de {len(partes)})"
            out.append(v)
    return out


# --------------------------------------------------------------------------
# HTML
# --------------------------------------------------------------------------
CSS = """
:root{--txt:#1d2733;--sec:#4d5b6a;--lnk:#0b5394;--bg:#fff;--ln:#d9dee4}
@media (prefers-color-scheme:dark){:root{--txt:#e3e7ec;--sec:#a6b1bd;--lnk:#8fbef0;--bg:#15191e;--ln:#333b44}}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--txt);font:17px/1.65 Georgia,"Times New Roman",serif}
main{max-width:46rem;margin:0 auto;padding:2rem 1.25rem 4rem}
nav,footer,.ctx{font:14px/1.5 system-ui,-apple-system,"Segoe UI",sans-serif;color:var(--sec)}
a{color:var(--lnk)}
h1{font-size:1.6rem;line-height:1.25;margin:.4rem 0 1.2rem}
.ctx{margin:0 0 .2rem}
p{margin:0 0 .8rem}
.tab{font:14px/1.5 system-ui,sans-serif;overflow-wrap:anywhere}
footer{border-top:1px solid var(--ln);margin-top:2.5rem;padding-top:1rem}
.pn{display:flex;justify-content:space-between;gap:1rem;margin-top:1.5rem}
ul{padding-left:1.2rem}li{margin:.2rem 0}
"""


def frase(s: str) -> str:
    """'DOS INSUMOS AGROPECUÁRIOS' -> 'Dos insumos agropecuários'."""
    if s and s.upper() == s:
        s = s.lower()
        return s[:1].upper() + s[1:]
    return s


def tema(u: dict) -> str:
    if u.get("nome"):
        return frase(u["nome"])
    for nv in reversed(NIVEIS):
        if nv in u["contexto"]:
            parte = u["contexto"][nv].split(" – ", 1)
            return frase(parte[1]) if len(parte) > 1 else parte[0]
    return ""


def resumo(u: dict) -> str:
    t = u["paras"][0] if u["paras"] else ""
    t = RE_ART.sub("", t).strip()
    return (t[:155] + "…") if len(t) > 156 else t


def link_oficial(fonte: dict) -> str:
    return fonte.get("url") or fonte.get("url_oficial", "")


def pagina(fonte: dict, u: dict, ant: dict | None, prox: dict | None) -> str:
    e = html.escape
    tm = tema(u)
    titulo = f"{fonte['sigla']}, {u['rotulo']}" + (f": {tm}" if tm else "")
    def _niv(t: str) -> str:
        a, _, b = t.partition(" – ")
        return f"{a} – {frase(b)}" if b else a
    ctx = " › ".join(_niv(u["contexto"][n]) for n in NIVEIS if n in u["contexto"])
    corpo = []
    for p in u["paras"]:
        cls = ' class="tab"' if " | " in p else ""
        corpo.append(f"<p{cls}>{e(p)}</p>")
    pn = '<div class="pn">'
    pn += f'<a href="{ant["slug"]}.html">‹ {e(ant["rotulo"])}</a>' if ant else "<span></span>"
    pn += f'<a href="{prox["slug"]}.html">{e(prox["rotulo"])} ›</a>' if prox else "<span></span>"
    pn += "</div>"
    return f"""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{e(titulo)}</title>
<meta name="description" content="{e(resumo(u))}">
<link rel="canonical" href="{e(BASE)}/{fonte['id']}/{u['slug']}.html">
<style>{CSS}</style></head>
<body><main>
<nav><a href="../index.html">Início</a> › <a href="index.html">{e(fonte['sigla'])}</a></nav>
{f'<p class="ctx">{e(ctx)}</p>' if ctx else ''}
<h1>{e(titulo)}</h1>
<article>
{chr(10).join(corpo)}
</article>
{pn}
<footer>Reprodução não oficial de {e(fonte['titulo'])}, gerada em {HOJE:%d/%m/%Y}.
Texto oficial: <a href="{e(link_oficial(fonte))}">{e(link_oficial(fonte))}</a>.
Trechos revogados ou com redação substituída foram omitidos; verifique sempre o texto oficial.</footer>
</main></body></html>
"""


def indice_fonte(fonte: dict, unidades: list[dict]) -> str:
    e = html.escape
    itens = "\n".join(
        f'<li><a href="{u["slug"]}.html">{e(u["rotulo"])}</a>'
        f'{": " + e(tema(u)) if tema(u) else ""}</li>' for u in unidades)
    return f"""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{e(fonte['titulo'])}: índice por artigo</title>
<style>{CSS}</style></head>
<body><main>
<nav><a href="../index.html">Início</a></nav>
<h1>{e(fonte['titulo'])}</h1>
<p>Índice com {len(unidades)} páginas. Texto oficial em
<a href="{e(link_oficial(fonte))}">{e(link_oficial(fonte))}</a>.</p>
<ul>
{itens}
</ul>
</main></body></html>
"""


def indice_geral(geradas: list[tuple[dict, int]]) -> str:
    e = html.escape
    itens = "\n".join(
        f'<li><a href="{f["id"]}/index.html">{e(f["titulo"])}</a> ({n} páginas)</li>'
        for f, n in geradas)
    return f"""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Reforma tributária do consumo: legislação por artigo (IBS, CBS, IS)</title>
<meta name="msvalidate.01" content="82CCDDA0F76C034A76A0118CC4ADAD71" />
<meta name="description" content="LC 214/2025, LC 227/2026, Decreto 12.955, Resoluções CGIBS e Nota Técnica, com uma página por artigo.">
<style>{CSS}</style></head>
<body><main>
<h1>Reforma tributária do consumo: legislação por artigo</h1>
<p>Reprodução não oficial, uma página por artigo, das normas do IBS, da CBS e do IS.
Atualizada em {HOJE:%d/%m/%Y}. Em caso de divergência, prevalece o texto oficial.</p>
<ul>
{itens}
</ul>
</main></body></html>
"""


# --------------------------------------------------------------------------
# Principal
# --------------------------------------------------------------------------
BASE = os.environ.get("SITE_BASE_URL", "https://example.github.io/repo").rstrip("/")


def main() -> int:
    cfg = json.loads((RAIZ / "fontes.json").read_text(encoding="utf-8"))
    SAIDA.mkdir(exist_ok=True)
    urls: list[str] = [f"{BASE}/index.html"]
    geradas: list[tuple[dict, int]] = []

    for fonte in cfg["fontes"]:
        raw = obter_bytes(fonte)
        if raw is None:
            continue
        try:
            if fonte["tipo"] == "planalto_html":
                unidades = segmentar(paragrafos_planalto(raw))
            elif fonte["tipo"] == "pdf_artigos":
                unidades = segmentar(paragrafos_pdf(raw))
            elif fonte["tipo"] == "xlsx_cclasstrib":
                unidades = unidades_cclasstrib(raw)
            elif fonte["tipo"] == "csv_indop":
                unidades = unidades_indop(raw)
            elif fonte["tipo"] == "xlsx_nbs":
                unidades = unidades_nbs(raw)
            else:
                unidades = segmentar_blocos(paragrafos_nt(raw))
        except Exception as ex:  # noqa: BLE001
            aviso(f"{fonte['id']}: erro ao processar ({ex})")
            continue
        unidades = dividir_grandes(unidades)
        if not unidades:
            aviso(f"{fonte['id']}: nenhum conteúdo extraído")
            continue

        # checagens de sanidade
        arts = [u for u in unidades if u["tipo"] == "artigo"]
        print(f"[{fonte['id']}] {len(unidades)} páginas ({len(arts)} de artigos)")
        amostra = [u for u in unidades if u["slug"] in ("art-138", "art-7-a", "anexo-ix")] or unidades[:2]
        det = "\n".join(
            f"--- {u['slug']} | tema: {tema(u)} | ctx: {' > '.join(u['contexto'].values())}\n"
            + "\n".join(p[:220] for p in u["paras"][:6]) for u in amostra)
        nota(f"{fonte['id']}: {len(unidades)} páginas, {len(arts)} artigos",
             f"primeira: {unidades[0]['slug']} | última: {unidades[-1]['slug']}\n"
             f"anexos: {', '.join(u['slug'] for u in unidades if u['tipo'] == 'anexo')}\n{det}")
        if fonte["id"] == "lc214":
            slugs = {u["slug"] for u in unidades}
            for obrig in ("art-138", "art-7-a", "anexo-ix"):
                if obrig not in slugs:
                    aviso(f"lc214: página esperada ausente: {obrig}")
            if len(arts) < 500:
                aviso(f"lc214: só {len(arts)} artigos extraídos (esperado > 500)")

        pasta = SAIDA / fonte["id"]
        pasta.mkdir(exist_ok=True)
        vistos: set[str] = set()
        for i, u in enumerate(unidades):
            base_slug, k = u["slug"], 2
            while u["slug"] in vistos:
                u["slug"] = f"{base_slug}-{k}"
                k += 1
            vistos.add(u["slug"])
        for i, u in enumerate(unidades):
            ant = unidades[i - 1] if i > 0 else None
            prox = unidades[i + 1] if i + 1 < len(unidades) else None
            (pasta / f"{u['slug']}.html").write_text(pagina(fonte, u, ant, prox), encoding="utf-8")
            urls.append(f"{BASE}/{fonte['id']}/{u['slug']}.html")
        (pasta / "index.html").write_text(indice_fonte(fonte, unidades), encoding="utf-8")
        urls.append(f"{BASE}/{fonte['id']}/index.html")
        geradas.append((fonte, len(unidades)))

    if not geradas:
        print("Nenhuma fonte processada.", file=sys.stderr)
        return 1

    (SAIDA / "index.html").write_text(indice_geral(geradas), encoding="utf-8")
    hoje = HOJE.isoformat()
    sm = ['<?xml version="1.0" encoding="UTF-8"?>',
          '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">']
    sm += [f"<url><loc>{html.escape(u)}</loc><lastmod>{hoje}</lastmod></url>" for u in urls]
    sm.append("</urlset>")
    (SAIDA / "sitemap.xml").write_text("\n".join(sm), encoding="utf-8")
    (SAIDA / "robots.txt").write_text(f"User-agent: *\nAllow: /\nSitemap: {BASE}/sitemap.xml\n")
    chave = cfg.get("indexnow_key")
    if chave:
        (SAIDA / f"{chave}.txt").write_text(chave)
    (SAIDA / ".nojekyll").write_text("")
    (SAIDA / "urls.txt").write_text("\n".join(urls))

    print(f"\n{len(urls)} URLs geradas em {SAIDA}")
    if avisos:
        print("\nAVISOS:\n- " + "\n- ".join(avisos))
    return 0


if __name__ == "__main__":
    sys.exit(main())
